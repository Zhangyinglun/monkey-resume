"""Behavioral coverage for versioned styles and bounded layout tuning."""
from __future__ import annotations

import copy
import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest import mock

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from scripts.extract_resume_text import extract_docx_text, extract_resume_blocks
from scripts.layout_auto_tuner import auto_fit_layout
from templates.docx_resume_template import generate_docx
from templates.layout_settings import DOCX_DEFAULT_SETTINGS
from templates.style_profiles import bounded_layouts, get_style
from tests.sample_data import sample_resume
from tests.test_pdf_pipeline import quality_report


class StyleProfileTests(unittest.TestCase):
    def test_legacy_retains_approved_geometry_and_headings(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            doc = Document(generate_docx("old.docx", sample_resume(), directory, style="legacy-v1"))
        self.assertEqual(doc.styles["Normal"].font.size.pt, 10.5)
        self.assertEqual(doc.paragraphs[0].runs[0].font.size.pt, 16.5)
        self.assertEqual(doc.paragraphs[2].text, "PROFESSIONAL EXPERIENCE")
        self.assertAlmostEqual(doc.sections[0].top_margin.mm, 12.7)
        self.assertEqual(get_style("legacy-v1").layout, DOCX_DEFAULT_SETTINGS)

    def test_new_style_is_black_left_aligned_and_has_no_layout_tables(self) -> None:
        resume = sample_resume()
        before = copy.deepcopy(resume)
        with tempfile.TemporaryDirectory() as directory:
            doc = Document(generate_docx("new.docx", resume, directory))
        self.assertEqual(resume, before)
        self.assertEqual(str(doc.styles["Normal"].font.color.rgb), "000000")
        self.assertEqual(doc.tables, [])
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, "")
        self.assertEqual(doc.sections[0].footer.paragraphs[0].text, "")
        for bullet in resume["experience"][0]["bullets"]:
            paragraph = next(p for p in doc.paragraphs if p.text == "• " + bullet)
            self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertIn("Skills", [p.text for p in doc.paragraphs])
        self.assertNotIn("Summary", [p.text for p in doc.paragraphs])

    def test_dates_anchor_to_right_margin_and_keep_reading_boundaries(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(generate_docx("dates.docx", sample_resume(), directory))
            doc = Document(path)
            row = next(p for p in doc.paragraphs if p.text.startswith("Example Corp"))
            tabs = list(row._p.iter(qn("w:ptab")))
            self.assertEqual(len(tabs), 1)
            self.assertEqual(tabs[0].get(qn("w:alignment")), "right")
            self.assertEqual(tabs[0].get(qn("w:relativeTo")), "margin")
            self.assertIn("Example Corp\t 2020–Present", extract_docx_text(path))
            self.assertTrue(any("Example Corp\t 2020–Present" in block["text"] for block in extract_resume_blocks(path)))

    def test_unreadable_override_fails_before_writing(self) -> None:
        profile = get_style()
        for layout in (
            replace(profile.layout, font_size_scale=0.9),
            replace(profile.layout, margin_top_mm=5.0),
            replace(profile.layout, font_size_scale=float("nan")),
            replace(profile.layout, item_spacing_scale=0.7),
        ):
            with self.subTest(layout=layout), tempfile.TemporaryDirectory() as directory:
                with self.assertRaises(ValueError):
                    generate_docx("bad.docx", sample_resume(), directory, layout)
                self.assertEqual(list(Path(directory).iterdir()), [])

    def test_compact_and_all_trials_respect_font_floor_and_caller_margins(self) -> None:
        profile = get_style()
        initial = replace(profile.layout, margin_side_inch=0.75)
        compact = profile.resolve_layout(replace(initial, compact_mode=True))
        for layout in [compact, *bounded_layouts(profile, initial)]:
            self.assertGreaterEqual(min(profile.font_sizes(layout).values()), 10.5)
            self.assertEqual(layout.margin_side_inch, 0.75)
            profile.validate(layout)
        self.assertEqual(bounded_layouts(profile, initial)[0].font_size_scale, initial.font_size_scale)

    def test_auto_fit_stops_at_first_passing_trial_without_changing_content(self) -> None:
        resume = sample_resume()
        before = copy.deepcopy(resume)
        reports = [quality_report(page_count=2), quality_report(page_count=2), quality_report()]
        with mock.patch("scripts.layout_auto_tuner.generate_resume", return_value="fake.pdf"), mock.patch(
            "scripts.layout_auto_tuner.check_pdf_file", side_effect=reports
        ):
            result = auto_fit_layout(resume, output_file="resume.pdf", max_trials=6)
        self.assertEqual(result.trials_run, 3)
        self.assertEqual(result.best_layout.font_size_scale, 1.0)
        self.assertEqual(resume, before)

    def test_tight_spacing_still_allows_font_reduction(self) -> None:
        profile = get_style()
        initial = replace(profile.layout, item_spacing_scale=0.75, section_spacing_scale=0.75, line_height_scale=0.7)
        candidates = bounded_layouts(profile, initial)
        self.assertTrue(any(layout.effective_font_size_scale < 1.0 for layout in candidates))
        def rendered_report(path: Path, *, minimum_font_pt: float) -> dict:
            layout = render.call_args.kwargs["layout"]
            return quality_report(page_count=1 if layout.effective_font_size_scale < 1.0 else 2)
        with mock.patch("scripts.layout_auto_tuner.generate_resume", return_value="fake.pdf") as render, mock.patch(
            "scripts.layout_auto_tuner.check_pdf_file", side_effect=rendered_report
        ):
            result = auto_fit_layout(sample_resume(), output_file="resume.pdf", max_trials=6, hint_layout=initial)
        self.assertEqual(result.best_report["verdict"], "PASS")
        self.assertEqual(profile.font_sizes(result.best_layout)["body_font_size"], 10.5)

    def test_auto_fit_does_not_claim_success_when_content_cannot_fit(self) -> None:
        with mock.patch("scripts.layout_auto_tuner.generate_resume", return_value="fake.pdf"), mock.patch(
            "scripts.layout_auto_tuner.check_pdf_file", return_value=quality_report(page_count=2)
        ):
            result = auto_fit_layout(sample_resume(), output_file="resume.pdf", max_trials=20)
        self.assertEqual(result.best_report["verdict"], "NEED-ADJUSTMENT")
        self.assertLessEqual(result.trials_run, 6)
        self.assertGreaterEqual(min(get_style().font_sizes(result.best_layout).values()), 10.5)

    def test_snapshot_is_json_serializable_and_records_actual_half_point_sizes(self) -> None:
        profile = get_style()
        snapshot = json.loads(json.dumps(profile.snapshot(profile.layout)))
        self.assertEqual(snapshot["actual_font_sizes_pt"]["header_font_size"], 19.0)
        self.assertEqual(snapshot["layout"]["margin_top_mm"], 15.24)
        self.assertEqual(snapshot["tokens"]["accent_color"], [0, 0, 0])

    def test_qa_uses_selected_font_threshold(self) -> None:
        self.assertEqual(quality_report(min_font_pt=10.0)["verdict"], "PASS")
        self.assertEqual(quality_report(min_font_pt=10.0, minimum_font_pt=10.5)["verdict"], "NEED-ADJUSTMENT")
        self.assertEqual(quality_report(min_font_pt=10.5, minimum_font_pt=10.5)["verdict"], "PASS")


if __name__ == "__main__":
    unittest.main()
