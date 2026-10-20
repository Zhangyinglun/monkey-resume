from __future__ import annotations

import copy
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document

from scripts.extract_resume_text import extract_docx_text
from scripts.layout_auto_tuner import auto_fit_layout
from templates.document_fonts import resolve_font_files
from templates.docx_resume_template import (
    export_pdf,
    generate_docx,
    generate_resume,
)
from templates.layout_settings import LayoutSettings
from templates.style_profiles import DEFAULT_STYLE, get_style
from tests.rendering_support import rendering_available
from tests.sample_data import sample_resume


class DocxTemplateTests(unittest.TestCase):
    def test_selected_education_facts_survive_docx_rendering(self) -> None:
        resume = sample_resume()
        resume["education"][0]["location"] = "Seattle"
        with tempfile.TemporaryDirectory() as directory:
            doc = Document(generate_docx("education.docx", resume, directory))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Example University | Seattle", text)
        self.assertIn(resume["education"][0]["degree"], text)
        self.assertIn(resume["education"][0]["dates"], text)

    def test_work_location_is_rendered_and_missing_location_blocks_output(self) -> None:
        resume = sample_resume()
        with tempfile.TemporaryDirectory() as directory:
            path = generate_docx("complete.docx", resume, directory)
            doc = Document(path)
            self.assertIn("Software Engineer | Seattle", [p.text for p in doc.paragraphs])
            resume["experience"][0].pop("location")
            with self.assertRaisesRegex(ValueError, r"experience\[0\].*location"):
                generate_docx("missing.docx", resume, directory)
            self.assertFalse((Path(directory) / "missing.docx").exists())

    def test_contact_is_immediately_followed_by_experience(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            doc = Document(generate_docx("resume.docx", sample_resume(), directory))
        self.assertEqual(doc.paragraphs[2].text, "Experience")
        self.assertNotIn("SUMMARY", [p.text for p in doc.paragraphs])
        self.assertTrue(all(p.text.strip() for p in doc.paragraphs[:4]))

    def test_old_summary_input_is_rejected_before_creating_artifacts(self) -> None:
        resume = sample_resume()
        resume["summary"] = ""
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, "summary"):
                generate_docx("resume.docx", resume, directory)
            self.assertEqual(list(Path(directory).iterdir()), [])

    def test_approved_readable_defaults(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            doc = Document(generate_docx("default.docx", sample_resume(), directory))
            self.assertEqual(doc.styles["Normal"].font.name, "Arial")
            self.assertEqual(doc.styles["Normal"].font.size.pt, 11.0)
            self.assertEqual(doc.paragraphs[0].runs[0].font.size.pt, 19.0)
            self.assertAlmostEqual(doc.sections[0].top_margin.mm, 15.24)
            self.assertAlmostEqual(doc.sections[0].bottom_margin.mm, 15.24)

    def test_explicit_layout_and_plain_text_roundtrip(self) -> None:
        resume = sample_resume()
        resume["experience"][0]["bullets"][0] = "Built R&D services with Java Map<K,V> and C++ & C#."
        before = copy.deepcopy(resume)
        layout = LayoutSettings(font_size_scale=1.0, margin_side_inch=0.75)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(generate_docx("resume.docx", resume, directory, layout, style=DEFAULT_STYLE))
            doc = Document(path)
            self.assertEqual(doc.styles["Normal"].font.size.pt, 11.0)
            self.assertAlmostEqual(doc.sections[0].left_margin.mm, 19.05)
            self.assertIn(resume["experience"][0]["bullets"][0], extract_docx_text(path))
        self.assertEqual(resume, before)

    def test_line_height_changes_spacing_not_content(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            texts, spacings = [], []
            for scale in (1.0, 0.95, 0.9, 0.88):
                path = Path(
                    generate_docx(
                        f"{scale}.docx",
                        sample_resume(),
                        directory,
                        LayoutSettings(line_height_scale=scale),
                        style=DEFAULT_STYLE,
                    )
                )
                texts.append(extract_docx_text(path))
                paragraph = next(
                    p for p in Document(path).paragraphs if p.text.startswith("• Built distributed")
                )
                spacings.append(paragraph.paragraph_format.line_spacing.pt)
            self.assertEqual(len(set(texts)), 1)
            self.assertGreater(spacings[0], spacings[1])
            self.assertTrue(all(a >= b for a, b in zip(spacings, spacings[1:])))

    def test_passing_default_is_not_squeezed_for_whitespace(self) -> None:
        before = sample_resume()
        with (
            mock.patch("scripts.layout_auto_tuner.generate_resume", return_value="fake.pdf"),
            mock.patch(
                "scripts.layout_auto_tuner.check_pdf_file",
                return_value={
                    "verdict": "PASS",
                    "checks": [
                        {"name": "bottom_margin", "passed": True, "detail": {"bottom_mm": 40}}
                    ],
                },
            ),
        ):
            result = auto_fit_layout(before, output_file="resume.pdf", max_trials=2)
        self.assertEqual(result.best_layout, get_style().layout)
        self.assertEqual(result.trials_run, 1)
        self.assertEqual(before, sample_resume())

    def test_missing_converter_has_descriptive_error(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(generate_docx("resume.docx", sample_resume(), directory))
            with mock.patch("templates.docx_resume_template.find_soffice", return_value=None):
                with self.assertRaises(RuntimeError):
                    export_pdf(path)

    def test_converter_timeout_does_not_publish_pdf(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(generate_docx("resume.docx", sample_resume(), directory, style=DEFAULT_STYLE))
            with (
                mock.patch("templates.docx_resume_template.find_soffice", return_value="mock-soffice"),
                mock.patch("templates.docx_resume_template.conversion_environment", return_value={}),
                mock.patch("templates.docx_resume_template.run_office", side_effect=subprocess.TimeoutExpired("soffice", 120)),
            ):
                with self.assertRaises(subprocess.TimeoutExpired):
                    export_pdf(path, style=DEFAULT_STYLE)
            self.assertFalse(path.with_suffix(".pdf").exists())

    def test_missing_selected_font_blocks_conversion(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(generate_docx("resume.docx", sample_resume(), directory))
            with (
                mock.patch("templates.docx_resume_template.find_soffice", return_value="mock-soffice"),
                mock.patch("templates.document_fonts.Path.is_dir", return_value=False),
                mock.patch("templates.document_fonts.shutil.which", return_value=None),
            ):
                with self.assertRaisesRegex(RuntimeError, "Arial Regular and Bold"):
                    export_pdf(path)
            self.assertFalse(path.with_suffix(".pdf").exists())

    def test_converter_substitution_is_rejected(self) -> None:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen.canvas import Canvas

        try:
            font_paths = resolve_font_files("Arial")
        except RuntimeError as exc:
            self.skipTest(str(exc))
        regular = next(
            path for path in font_paths
            if TTFont("inspect", str(path)).face.styleName == b"Regular"
        )
        pdfmetrics.registerFont(TTFont("ConverterRegularOnly", str(regular)))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(generate_docx("resume.docx", sample_resume(), directory, style=DEFAULT_STYLE))
            def converter(command: list[str], **kwargs: object) -> subprocess.CompletedProcess:
                canvas = Canvas(str(path.with_suffix(".pdf")))
                canvas.setFont(replacement, 11)
                canvas.drawString(50, 750, "This converter lost the requested font or bold face.")
                canvas.save()
                return subprocess.CompletedProcess(command, 0, "", "")

            for replacement in ("Helvetica", "ConverterRegularOnly"):
                path.with_suffix(".pdf").unlink(missing_ok=True)
                with (
                    self.subTest(replacement=replacement),
                    mock.patch("templates.docx_resume_template.find_soffice", return_value="mock-soffice"),
                    mock.patch("templates.docx_resume_template.run_office", side_effect=converter),
                ):
                    with self.assertRaisesRegex(RuntimeError, "Refusing font substitution"):
                        export_pdf(path, style=DEFAULT_STYLE)


@unittest.skipUnless(
    rendering_available(), "LibreOffice, Poppler and Arial required"
)
class DocxRenderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.render_style = DEFAULT_STYLE
        if not rendering_available():
            raise unittest.SkipTest("LibreOffice, Poppler and Arial required")

    def test_actual_docx_contact_wrapping_blocks_pdf_qa(self) -> None:
        from scripts.check_pdf_quality import check_pdf_file

        resume = sample_resume()
        resume["contact"] = "Seattle | alex@example.org | linkedin.com/in/example | " * 4
        with tempfile.TemporaryDirectory() as directory:
            pdf = Path(generate_resume("wrapped-contact.pdf", resume, directory))
            report = check_pdf_file(pdf)
        check = next(c for c in report["checks"] if c["name"] == "contact_single_line")
        self.assertTrue(check["detail"]["measured"])
        self.assertGreater(check["detail"]["line_count"], 1)
        self.assertFalse(check["passed"])
        self.assertEqual(report["verdict"], "NEED-ADJUSTMENT")

    def test_same_source_pdf_matches_repeated_pdf_render(self) -> None:
        from pdf2image import convert_from_path
        from PIL import Image

        with tempfile.TemporaryDirectory(prefix="resume space ") as directory:
            root = Path(directory)
            pdf = Path(generate_resume("候选 Resume.pdf", sample_resume(), root, style=self.render_style))
            copied = root / "independent.docx"
            shutil.copy2(pdf.with_suffix(".docx"), copied)
            script = Path(__file__).resolve().parents[1] / "scripts/render_pdf_reference.py"
            subprocess.run(
                [
                    sys.executable,
                    str(script),
                    str(copied),
                    "--output_dir",
                    str(root / "preview"),
                    "--dpi",
                    "96",
                    "--style",
                    self.render_style,
                ],
                check=True,
                capture_output=True,
                timeout=120,
            )
            expected = convert_from_path(pdf, dpi=96)
            pages = sorted((root / "preview").glob("page-*.png"))
            self.assertEqual(len(expected), 1)
            self.assertEqual(len(pages), 1)
            with Image.open(pages[0]) as actual:
                self.assertEqual(expected[0].size, actual.size)
                self.assertEqual(expected[0].tobytes(), actual.tobytes())

    def test_default_render_uses_actual_arial_fonts(self) -> None:
        import pdfplumber

        resolve_font_files("Arial")
        with tempfile.TemporaryDirectory() as directory:
            pdf = Path(generate_resume("fonts.pdf", sample_resume(), directory))
            with pdfplumber.open(pdf) as rendered:
                fonts = {char["fontname"].split("+")[-1] for page in rendered.pages for char in page.chars}
        self.assertIn("ArialMT", fonts)
        self.assertIn("Arial-BoldMT", fonts)
        self.assertTrue(all(name.startswith("Arial") for name in fonts), fonts)

    def test_real_overflow_detected_without_editing_content(self) -> None:
        from scripts.check_pdf_quality import check_pdf_file

        resume = sample_resume()
        resume["experience"][0]["bullets"] *= 15
        before = copy.deepcopy(resume)
        with tempfile.TemporaryDirectory() as directory:
            pdf = Path(generate_resume("overflow.pdf", resume, directory, style=self.render_style))
            report = check_pdf_file(pdf)
        count = next(c for c in report["checks"] if c["name"] == "page_count")
        self.assertGreater(count["detail"]["count"], 1)
        self.assertEqual(report["verdict"], "NEED-ADJUSTMENT")
        self.assertEqual(resume, before)


if __name__ == "__main__":
    unittest.main()
