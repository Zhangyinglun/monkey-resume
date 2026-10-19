"""Editable Word adapter using python-docx; PDF is exported from this same document.

This is a template adapter, not a pagination or PDF engine. Layout is delegated
entirely to WordprocessingML and LibreOffice. See references/docx-output.md.
"""
from __future__ import annotations

import hashlib
import platform
import re
import shutil
import subprocess
import tempfile
from importlib.metadata import version
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Mm, Pt, RGBColor
from docx.text.run import Run
from reportlab.pdfbase.pdfmetrics import stringWidth

from scripts.resume_shared import normalize_skill_items, validate_resume_content
from templates.design_tokens import DesignTokens
from templates.document_fonts import conversion_environment, resolve_font_files, validate_pdf_fonts
from templates.layout_settings import LayoutSettings
from templates.style_profiles import DEFAULT_STYLE, get_style

PROJECT_ROOT = Path(__file__).resolve().parent.parent


def _explicit_fonts(doc: Any, family: str) -> None:
    """Bind portable output text to its declared family rather than reader themes."""
    for style in doc.styles:
        if style.type in (1, 2, 3):
            style.font.name = family
    for root in (doc.styles.element, doc.element):
        for run in root.iter(qn("w:r")):
            props = run.find(qn("w:rPr"))
            if props is None:
                props = OxmlElement("w:rPr")
                run.insert(0, props)
            if props.find(qn("w:rFonts")) is None:
                props.insert(0, OxmlElement("w:rFonts"))
        for fonts in root.iter(qn("w:rFonts")):
            for attr in list(fonts.attrib):
                if attr.lower().endswith("theme"):
                    del fonts.attrib[attr]
            for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + slot), family)


def _add_text(paragraph: Any, text: str, *, bold: bool = False) -> None:
    """Keep contact URLs clickable without interpreting candidate text as markup."""
    pattern = r"(?:https?://[^\s|]+|(?:www\.|linkedin\.com/|github\.com/)[^\s|]+|[\w.+-]+@[\w.-]+\.[A-Za-z]{2,})"
    cursor = 0
    for match in re.finditer(pattern, text):
        paragraph.add_run(text[cursor:match.start()]).bold = bold
        label = match.group()
        url = label if label.startswith(('http://', 'https://')) else (
            'mailto:' + label if '@' in label else 'https://' + label
        )
        relationship = paragraph.part.relate_to(
            url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
            is_external=True,
        )
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), relationship)
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        if bold:
            props.append(OxmlElement('w:b'))
        run.append(props)
        value = OxmlElement('w:t')
        value.text = label
        run.append(value)
        link.append(run)
        paragraph._p.append(link)
        cursor = match.end()
    paragraph.add_run(text[cursor:]).bold = bold


def generate_docx(output_file: str, content_dict: dict[str, Any], base_dir: str,
                  layout: LayoutSettings | None = None,
                  tokens: DesignTokens | None = None, *, style: str = DEFAULT_STYLE) -> str:
    validate_resume_content(content_dict)
    profile = get_style(style)
    settings, design_tokens = profile.resolve_layout(layout), tokens or profile.tokens
    profile.validate(settings, design_tokens)
    target_dir = Path(base_dir).expanduser().resolve()
    if target_dir == PROJECT_ROOT or PROJECT_ROOT in target_dir.parents:
        raise ValueError('Output must live outside the Skill package.')
    if Path(output_file).name != output_file or Path(output_file).suffix.lower() != '.docx':
        raise ValueError('output_file must be a .docx filename only.')
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / output_file
    if target.exists():
        raise FileExistsError(f'Refusing to overwrite existing DOCX: {target}')
    font_scale, line_height_scale = settings.effective_font_size_scale, settings.effective_line_height_scale
    section_spacing_scale, item_spacing_scale = settings.effective_section_spacing_scale, settings.effective_item_spacing_scale
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(settings.margin_top_mm), Mm(settings.margin_bottom_mm)
    section.left_margin = section.right_margin = Inches(settings.margin_side_inch)
    normal = doc.styles['Normal']
    normal.font.name = profile.font_family
    normal.font.size = Pt(design_tokens.body_font_size * font_scale)
    normal.font.color.rgb = RGBColor(*design_tokens.body_ink_color)
    normal.paragraph_format.space_after = Pt(design_tokens.body_space_after * item_spacing_scale)
    normal.paragraph_format.line_spacing = Pt(max(design_tokens.body_leading * line_height_scale, design_tokens.body_font_size * font_scale * 1.1))

    def paragraph(text: str, size: float = design_tokens.body_font_size, *, bold: bool = False,
                  after: float = design_tokens.body_space_after, center: bool = False,
                  keep: bool = False, leading: float = design_tokens.body_leading) -> Any:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after * item_spacing_scale)
        p.paragraph_format.keep_with_next = keep
        p.paragraph_format.widow_control = True
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _add_text(p, text, bold=bold)
        # Compress leading independently of font size, with a glyph-clearance floor.
        p.paragraph_format.line_spacing = Pt(max(leading * line_height_scale, size * font_scale * 1.1))
        for element in p._p.iter(qn("w:r")):
            Run(element, p).font.size = Pt(size * font_scale)
        return p

    def heading(title: str, *, literal: bool = False) -> None:
        p = paragraph(title if literal else profile.heading(title), design_tokens.section_font_size, bold=True, after=design_tokens.section_space_after,
                      keep=True, leading=design_tokens.section_leading)
        p.paragraph_format.space_before = Pt(design_tokens.section_space_before * section_spacing_scale)
        for run in p.runs:
            run.font.color.rgb = RGBColor(*design_tokens.accent_color)
        borders = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        for key, value in {'val': 'single', 'sz': str(round(design_tokens.section_hr_thickness * 8)),
                           'color': ''.join(f'{c:02X}' for c in design_tokens.accent_color), 'space': '3'}.items():
            bottom.set(qn('w:' + key), value)
        borders.append(bottom)
        p._p.get_or_add_pPr().append(borders)

    def row(left: str, dates: str) -> None:
        follows_bullet = bool(doc.paragraphs and doc.paragraphs[-1].text.startswith("• "))
        p = paragraph(left, design_tokens.company_font_size, bold=True, after=design_tokens.company_space_after,
                      keep=True, leading=design_tokens.company_leading)
        if follows_bullet:
            p.paragraph_format.space_before = Pt(design_tokens.company_space_before * section_spacing_scale)
        width = section.page_width - section.left_margin - section.right_margin
        if dates:
            # Helvetica's Arial-compatible widths plus clearance give long rows
            # a separate date line; actual rendering is still checked in QA.
            text_width = stringWidth(left, "Helvetica-Bold", design_tokens.company_font_size * font_scale)
            date_width = stringWidth(dates, "Helvetica", design_tokens.dates_font_size * font_scale)
            if text_width + date_width + 24 > Emu(width).pt:
                p.paragraph_format.space_after = Pt(0)
                date = paragraph(dates, design_tokens.dates_font_size,
                                 after=design_tokens.company_space_after, keep=True)
                date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                return
        p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        date_run = p.add_run()
        date_run.add_tab()
        date_run.add_text(dates)
        date_run.bold = False
        date_run.font.size = Pt(design_tokens.dates_font_size * font_scale)

    def add_bullet(text: str) -> None:
        p = paragraph('• ' + text, design_tokens.bullet_font_size,
                      after=design_tokens.bullet_space_after)
        p.paragraph_format.left_indent = Pt(design_tokens.bullet_left_indent)
        p.paragraph_format.first_line_indent = Pt(-10)

    paragraph(content_dict['name'], design_tokens.header_font_size, bold=True, center=True,
              after=design_tokens.header_space_after, keep=True, leading=design_tokens.header_leading)
    paragraph(content_dict['contact'], design_tokens.contact_font_size, center=True,
              after=design_tokens.contact_space_after, keep=True)
    heading('PROFESSIONAL EXPERIENCE')
    for item in content_dict['experience']:
        row(item.get('company', ''), item.get('dates', ''))
        paragraph(' | '.join(str(item.get(k, '')) for k in ('title', 'location') if item.get(k)),
                  design_tokens.job_detail_font_size, after=design_tokens.job_detail_space_after, keep=True)
        for bullet in item.get('bullets', []):
            add_bullet(bullet)
    for key, title, extra in [('projects', 'PROJECTS', 'tech')]:
        if content_dict.get(key):
            heading(content_dict.get('project_heading', title), literal='project_heading' in content_dict)
            for item in content_dict[key]:
                row(' | '.join(str(item[k]) for k in ('name', extra) if item.get(k)), item.get('dates', ''))
                for bullet in item.get('bullets', []):
                    add_bullet(bullet)
    heading('TECHNICAL SKILLS')
    for skill in content_dict['skills']:
        p = paragraph('')
        _add_text(p, skill.get('category', '') + ': ', bold=True)
        _add_text(p, ', '.join(normalize_skill_items(skill.get('items', []))))
    for key, title, extra in [('certifications', 'CERTIFICATIONS', 'issuer'), ('awards', 'AWARDS', 'organization')]:
        if content_dict.get(key):
            heading(title)
            for item in content_dict[key]:
                row(' - '.join(str(item[k]) for k in ('name', extra) if item.get(k)), item.get('dates', ''))
    heading('EDUCATION')
    for edu in content_dict['education']:
        row(edu.get('school', ''), edu.get('dates', ''))
        if edu.get('degree'):
            paragraph(edu['degree'], design_tokens.education_font_size, after=design_tokens.education_space_after)
    _explicit_fonts(doc, profile.font_family)
    doc.save(target)
    return str(target)


def export_pdf(docx_path: Path, *, style: str = DEFAULT_STYLE) -> Path:
    """Use the installed mature office renderer; never silently fall back to another layout."""
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        raise RuntimeError('LibreOffice is required for DOCX layout QA and PDF export.')
    family = get_style(style).font_family
    target = docx_path.with_suffix('.pdf')
    if target.exists():
        raise FileExistsError(f'Refusing stale PDF export: {target}')
    with tempfile.TemporaryDirectory(prefix='resume-office-profile-') as profile:
        env = conversion_environment(family, Path(profile) / "fonts")
        result = subprocess.run(
            [soffice, '-env:UserInstallation=' + Path(profile).as_uri(), '--headless',
             '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(docx_path.parent), str(docx_path)],
            capture_output=True, text=True, timeout=120, check=False, env=env,
        )
    if result.returncode or not target.is_file() or not target.read_bytes().startswith(b'%PDF'):
        raise RuntimeError('LibreOffice PDF export failed: ' + result.stdout + result.stderr)
    validate_pdf_fonts(target, family)
    return target


def generate_resume(output_file: str, content_dict: dict[str, Any], base_dir: str = 'resume_output',
                    layout: LayoutSettings | None = None, *, style: str = DEFAULT_STYLE) -> str:
    """Render a DOCX and export its same-source PDF for QA and layout tuning."""
    if Path(output_file).suffix.lower() != '.pdf':
        raise ValueError('QA output_file must use .pdf.')
    source = generate_docx(str(Path(output_file).with_suffix('.docx')), content_dict, base_dir, layout, style=style)
    return str(export_pdf(Path(source), style=style))


def rendering_environment(*, style: str = DEFAULT_STYLE) -> dict[str, Any]:
    """Record the renderer and installed font files, without resume content."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    renderer_version = None
    if soffice:
        try:
            result = subprocess.run([soffice, "--version"], capture_output=True, text=True, timeout=15, check=False)
            if result.returncode == 0:
                renderer_version = result.stdout.strip() or None
        except (OSError, subprocess.SubprocessError):
            pass
    family = get_style(style).font_family
    try:
        font_files = resolve_font_files(family)
    except (OSError, RuntimeError, subprocess.SubprocessError):
        font_files = ()
    return {
        "platform": platform.platform(), "python_docx": version("python-docx"),
        "renderer": "LibreOffice Writer", "renderer_version": renderer_version,
        "font_family": family,
        "font_files": [{"name": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for path in sorted(font_files)],
    }
